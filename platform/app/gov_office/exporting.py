"""政务公文导出：审核通过的 DOCX 清洁稿与可追溯证据包。"""

from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from datetime import UTC, datetime
from typing import Any

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from sqlalchemy.orm import Session

from app.common.errors import BusinessError
from app.gov_office.documents import GovDocumentService
from app.gov_office.security import append_gov_audit
from core.abstraction.dtos import Actor

SENSITIVE_PATTERNS = [
    re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)"),
    re.compile(r"(?<!\d)\d{17}[\dXx](?!\d)"),
]


def _set_font(run, name: str, size: int) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _add_paragraph(doc: DocxDocument, text: str, *, size: int = 14, bold: bool = False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(text)
    _set_font(run, "仿宋_GB2312", size)
    run.bold = bold
    return paragraph


def _build_docx(draft: dict[str, Any]) -> bytes:
    doc = DocxDocument()
    section = doc.sections[0]
    section.top_margin = Pt(72)
    section.bottom_margin = Pt(72)
    section.left_margin = Pt(85)
    section.right_margin = Pt(85)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(draft["title"])
    _set_font(title_run, "方正小标宋简体", 22)
    title_run.bold = True

    if draft.get("recipient"):
        _add_paragraph(doc, draft["recipient"])

    for section_data in draft.get("sections") or []:
        heading = str(section_data.get("heading", "")).strip()
        if heading:
            heading_paragraph = _add_paragraph(doc, heading, size=16, bold=True)
            heading_paragraph.paragraph_format.line_spacing = 1.5
        content = str(section_data.get("content", "")).strip()
        if content:
            _add_paragraph(doc, content)

    doc.add_paragraph()
    authority = _add_paragraph(doc, draft["issuing_authority"])
    authority.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    date_line = _add_paragraph(doc, draft["document_date"])
    date_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


class GovDocumentExportService:
    def __init__(self, db: Session):
        self.db = db

    def export(self, draft_id: str, actor: Actor) -> tuple[bytes, str]:
        draft = GovDocumentService(self.db).get(draft_id, actor)
        if draft["status"] != "approved":
            raise BusinessError("DOCUMENT_EXPORT_NOT_APPROVED", "仅审核通过的公文可导出", 409)
        errors = [
            item for item in draft.get("checks", []) if item.get("severity") == "error"
        ]
        if errors:
            raise BusinessError("DOCUMENT_EXPORT_BLOCKED", "公文仍有阻断项，不能导出", 409)
        body = "\n".join(
            [
                draft["title"],
                draft.get("recipient", ""),
                *[
                    f"{section.get('heading', '')}\n{section.get('content', '')}"
                    for section in draft.get("sections", [])
                ],
            ]
        )
        if any(pattern.search(body) for pattern in SENSITIVE_PATTERNS):
            raise BusinessError(
                "DOCUMENT_EXPORT_BLOCKED", "正文含疑似手机号或身份证号，禁止导出", 409
            )
        docx_bytes = _build_docx(draft)
        docx_hash = hashlib.sha256(docx_bytes).hexdigest()
        exported_at = datetime.now(UTC).isoformat()
        append_gov_audit(
            self.db,
            actor,
            action="gov.document.export",
            resource=draft_id,
            detail={
                "docx_sha256": docx_hash,
                "draft_version": draft["version"],
                "content_hash": draft.get("content_hash", ""),
            },
        )
        self.db.commit()
        evidence = {
            "draft": draft,
            "export": {
                "exported_at": exported_at,
                "exported_by": actor.user_id,
                "docx_sha256": docx_hash,
                "notice": "本文件为审核通过的草稿导出，不代表签发、用印或正式发布。",
            },
        }
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
            bundle.writestr(f"{draft_id}.docx", docx_bytes)
            bundle.writestr(
                f"{draft_id}-evidence.json",
                json.dumps(evidence, ensure_ascii=False, indent=2),
            )
            bundle.writestr(
                "README.txt",
                "DOCX 为清洁稿；JSON 为事实、引用、审核与版本证据包。"
                "导出不代表签发、用印或正式发布。",
            )
        return archive.getvalue(), docx_hash
